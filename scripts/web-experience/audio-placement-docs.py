# 音声配置一覧を、Markdown と Excel に書き出す（元データ：out/audio-placement.json）。
#   npx tsx scripts/web-experience/audio-placement.ts && python3 scripts/web-experience/audio-placement-docs.py
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
d = json.load(open(ROOT / 'out' / 'audio-placement.json'))
rows, files = d['rows'], d['files']
OUT = ROOT / 'docs' / 'web-experience'

COLS = [
    ('lineId', 'セリフID', 16), ('phase', 'フェーズ', 18), ('scene', '場面', 24), ('role', '種類', 10), ('assetId', '映像（クリップ）', 16),
    ('speaker', '話者', 12), ('voice', '声', 7), ('wipe', 'ワイプ', 16), ('text', 'セリフ（字幕）', 60),
    ('sourceFile', '受領ファイル', 22), ('sourceFrom', '受領ファイル内 開始（秒）', 10), ('sourceTo', '同 終了（秒）', 10),
    ('seconds', '音声の長さ（秒）', 9), ('clipFrom', '映像内 開始（秒）', 9), ('clipTo', '映像内 終了（秒）', 9), ('subtitleTo', '字幕 終了（秒）', 9), ('clipSeconds', '映像の長さ（秒）', 9), ('check', '確認', 14),
]

ng = [r for r in rows if r['issues']]
md = ['# 13　音声配置一覧（音読さん Beta・16本 → 137文）', '',
      '> `npx tsx scripts/web-experience/audio-placement.ts && python3 scripts/web-experience/audio-placement-docs.py` で自動生成。Excel 版：`13-audio-placement.xlsx`。', '',
      '## 確認結果', '',
      '| 確認項目 | 方法 | 結果 |', '|---|---|---|',
      f'| 未使用 | 受領ファイルの「話している部分」が、どの文の区切りにも入っていないところがないか | {"なし" if not any(f["unused"] for f in files) else "あり"}（16本すべて、全体を使用） |',
      f'| 未配置 | 137文すべてに1文ずつの音声があり、映像に入っているか | {137 - sum(1 for r in rows if any("未配置" in i or "未使用" in i for i in r["issues"]))}/137 |',
      f'| 誤配置 | 受領ファイルの話者と、セリフの話者が一致するか（ファイル名で照合） | {"不一致なし" if not any("誤配置" in i for r in rows for i in r["issues"]) else "不一致あり"} |',
      f'| 重複再生 | 同じ文が2つの映像に入っていないか／映像の中で前の文と重ならないか | {"なし" if not any("重複" in i for r in rows for i in r["issues"]) else "あり"} |',
      f'| 音切れ | 1文の頭・お尻の音量と、受領ファイルで区切りの外側が無音か／映像の終わりより長くないか | {"なし" if not any("音切れ" in i for r in rows for i in r["issues"]) else "あり"} |',
      f'| 字幕の不一致 | 字幕は音声と同じ文から作る（1対1）。字幕が音声より先に消えないか | {"なし" if not any("字幕" in i for r in rows for i in r["issues"]) else "あり"} |',
      '', '完成した動画での照合（位置・欠落・重複・途中切れ）は `scripts/experience-verify.py` の結果を 12-open-issues.md に記載。', '',
      '## 受領ファイル（16本）', '', '| ファイル | 長さ（秒） | 文の数 | どの文にも入っていない音 |', '|---|---|---|---|']
for f in files:
    md.append(f'| {f["file"]} | {f["seconds"]} | {f["lines"]} | {"、".join(f["unused"]) or "なし"} |')
md += ['', '`_part1` と `_part2` は番号順に、0.9秒の間をあけてつないでから文に分けた。', '', '## 全137文の配置（再生順）', '',
       '| セリフID | フェーズ | 場面 | 種類 | 話者（声） | ワイプ | セリフ | 受領ファイル（秒） | 映像（開始秒） |', '|---|---|---|---|---|---|---|---|---|']
for r in rows:
    md.append(f'| `{r["lineId"]}` | {r["phase"]} | {r["scene"]} | {r["role"]} | {r["speaker"]}（{r["voice"]}） | {r["wipe"]} | {r["text"].replace("|", "｜")} | {r["sourceFile"]} {r["sourceFrom"]}〜{r["sourceTo"]} | {r["assetId"]} {r["clipFrom"]} |')
(OUT / '13-audio-placement.md').write_text('\n'.join(md) + '\n')

wb = Workbook()
ws = wb.active
ws.title = '音声配置一覧'
ws.append([c[1] for c in COLS])
for r in rows:
    ws.append([('OK' if not r['issues'] else '／'.join(r['issues'])) if k == 'check' else r.get(k) for k, _, _ in COLS])
head = PatternFill('solid', fgColor='0D0E10')
for i, (_, _, w) in enumerate(COLS, 1):
    c = ws.cell(row=1, column=i)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = head
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'B2'
ws.auto_filter.ref = ws.dimensions
ws2 = wb.create_sheet('受領ファイル')
ws2.append(['ファイル', '長さ（秒）', '文の数', 'どの文にも入っていない音'])
for f in files:
    ws2.append([f['file'], f['seconds'], f['lines'], '、'.join(f['unused']) or 'なし'])
for i, w in enumerate([26, 12, 10, 40], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.cell(row=1, column=i).font = Font(bold=True, color='FFFFFF')
    ws2.cell(row=1, column=i).fill = head
wb.save(OUT / '13-audio-placement.xlsx')
print(f'13-audio-placement.md / .xlsx：{len(rows)}文・確認で引っかかった文 {len(ng)}')
