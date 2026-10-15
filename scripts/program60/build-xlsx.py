"""60分版の設計書（分析・構成・講師用台本・素材管理・品質確認・アンケート・要確認事項）を Excel に書き出す。

使い方：
  node scripts/program60/build-pptx.mjs          # 先にスライドを作る（画像枠の位置を out/ に書き出す）
  node -e "import('./scripts/program60/content.mjs').then(m=>require('fs').writeFileSync('out/program60-content.json',JSON.stringify(m)))"
  python3 scripts/program60/build-xlsx.py
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
C = json.loads((ROOT / 'out' / 'program60-content.json').read_text())
PLACED = json.loads((ROOT / 'out' / 'program60-placements.json').read_text())
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'docs' / 'program60' / 'program60-design.xlsx'

FONT = 'Meiryo UI'
RED, INK, SOFT, GRAY, PALE, YELLOW = 'E60012', '1E2126', 'FDECEE', '5A6470', 'F3F4F6', 'FFF2CC'
thin = Side(style='thin', color='D5D9DE')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

wb = Workbook()
wb.remove(wb.active)


def sheet(name, title, lead):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws['A1'] = title
    ws['A1'].font = Font(name=FONT, size=16, bold=True, color=INK)
    ws['A2'] = lead
    ws['A2'].font = Font(name=FONT, size=10, color=GRAY)
    return ws


def table(ws, row, headers, rows, widths, fills=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=(fills or {}).get(j, INK))
        c.alignment = CENTER
        c.border = BORDER
    for i, r in enumerate(rows, row + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FONT, size=10, color=INK)
            c.alignment = WRAP
            c.border = BORDER
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    return row + len(rows)


phase_of = {p['key']: p for p in C['PHASES']}
media_of = {m['id']: m for m in C['MEDIA']}

# ── 0 概要 ──
ws = sheet('0_概要', 'コンサルティング営業体験（60分版）設計書', '大学生向け新卒採用プログラム。スライドは program60-slides.pptx（本編45枚＋補足6枚）')
rows = [
    ['①既存69枚の分析', '既存スライドごとに、残す・統合・移動・削除・変更の判断と、新しいスライド番号'],
    ['②全体構成', '6フェーズのタイムライン、スライド構成、参加型ワークの配置'],
    ['③改善後の本編スライド', 'program60-slides.pptx（別ファイル）。各スライドのノートに講師用台本の要約'],
    ['④講師用進行台本', 'スライドごとの説明・問いかけ・想定回答・フィードバック・ヒント・つなぎ・時間超過時の調整'],
    ['⑤画像・動画差し替え管理', '差し替え枠の番号・配置・目的・枠の大きさと比率・仮素材・状況'],
    ['⑥品質確認', '指示書20章の8項目のチェック結果'],
    ['⑦アンケート', '事前・事後の設問と、成果を見る指標'],
    ['⑧要確認事項', '推測で埋めずに残した事項。だれに確認するか'],
]
table(ws, 4, ['シート', '内容'], rows, [28, 100])
ws['A14'] = '黄色のセルは、運営側が記入する欄です（担当・差し替え日・状況など）。'
ws['A14'].font = Font(name=FONT, size=10, color=GRAY)
ws['A14'].fill = PatternFill('solid', fgColor=YELLOW)

# ── ① 既存69枚の分析 ──
ws = sheet('1_既存69枚の分析', '成果物①　既存資料（69枚）の改善分析', '判断：残す（改善）／統合／移動／削除／変更。移動先が「講師台本」のものは、画面から外して講師が使う')
last = table(ws, 4, ['既存No', '既存スライド', '判断', '新しいスライド', '理由・変えること'], C['OLD'], [8, 34, 14, 22, 70], {3: RED})
r = last + 2
ws.cell(row=r, column=1, value='判断ごとの枚数').font = Font(name=FONT, size=11, bold=True)
for k, lab in enumerate(['残す', '統合', '移動', '変更', '削除']):
    ws.cell(row=r + 1 + k, column=2, value=lab).font = Font(name=FONT, size=10)
    c = ws.cell(row=r + 1 + k, column=3, value=f'=COUNTIF($C$5:$C${last},"{lab}*")')
    c.font = Font(name=FONT, size=10, bold=True)
ws.cell(row=r + 6, column=2, value='合計').font = Font(name=FONT, size=10, bold=True)
ws.cell(row=r + 6, column=3, value=f'=SUM(C{r + 1}:C{r + 5})').font = Font(name=FONT, size=10, bold=True)
ws.cell(row=r + 8, column=1, value='新しく追加したスライド').font = Font(name=FONT, size=11, bold=True)
table_start = r + 9
for j, h in enumerate(['新しいスライド', '内容', 'ねらい'], 1):
    c = ws.cell(row=table_start, column=j + 1, value=h)
    c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor=INK)
    c.border = BORDER
for i, row in enumerate(C['NEW_ONLY'], table_start + 1):
    for j, v in enumerate(row, 2):
        c = ws.cell(row=i, column=j, value=v)
        c.font = Font(name=FONT, size=10)
        c.alignment = WRAP
        c.border = BORDER

# ── ④ 講師用進行台本（②より先に作り、②から参照する）──
ws4 = sheet('4_講師用進行台本', '成果物④　講師用進行台本', '開始・終了は「所要（分）」から自動計算。所要を変えると全体の時間が更新される。黄色＝当日の担当メモ')
hdr = ['No', 'フェーズ', '開始', '所要（分）', '終了', '進め方', '画面に表示する内容', '講師の説明・発話', '学生への問いかけ', '想定される学生の回答', '回答へのフィードバック', '答えが出ないときのヒント', '使用する画像・動画', '次へのつなぎ', '時間超過時の調整', '既存スライド', '当日メモ']
rows = []
for s in C['SLIDES']:
    n = s['n']
    media = '\n'.join(f"{m['id']}（{media_of[m['id']]['use']}）" for m in s.get('media', []))
    rows.append([s['id'], s['phase'], None, s['dur'], None, s['mode'], n['screen'], n['talk'], n['ask'], n['expect'], n['feedback'], n['hint'], media, n['next'], n['skip'], s['from'], ''])
last4 = table(ws4, 4, hdr, rows, [6, 12, 7, 8, 7, 11, 30, 48, 30, 34, 38, 30, 30, 26, 30, 16, 16], {4: RED})
for i in range(5, last4 + 1):
    ws4.cell(row=i, column=3, value='=0' if i == 5 else f'=E{i - 1}')
    ws4.cell(row=i, column=5, value=f'=C{i}+D{i}')
    for col in (3, 5):
        ws4.cell(row=i, column=col).number_format = '0.00'
        ws4.cell(row=i, column=col).alignment = CENTER
    ws4.cell(row=i, column=4).alignment = CENTER
    ws4.cell(row=i, column=4).font = Font(name=FONT, size=10, color='0000FF', bold=True)
    ws4.cell(row=i, column=17).fill = PatternFill('solid', fgColor=YELLOW)
tot = last4 + 1
ws4.cell(row=tot, column=2, value='合計').font = Font(name=FONT, size=10, bold=True)
ws4.cell(row=tot, column=4, value=f'=SUM(D5:D{last4})').font = Font(name=FONT, size=10, bold=True)
ws4.cell(row=tot, column=6, value=f'=IF(D{tot}<=60,"60分以内","60分を超えています")').font = Font(name=FONT, size=10, bold=True, color=RED)
ws4.conditional_formatting.add(f'A5:Q{last4}', FormulaRule(formula=['OR($F5="個人→投票",$F5="個人→ペア",$F5="グループ",$F5="個人→グループ",$F5="グループ発表",$F5="ペア",$F5="個人")'], fill=PatternFill('solid', fgColor='FFF7F7')))
ws4['A3'] = '薄い赤の行＝学生が手を動かす場面'
ws4['A3'].font = Font(name=FONT, size=9, color=RED)

# ── ② 全体構成 ──
ws = sheet('2_全体構成', '成果物②　新しい全体構成（60分）', '時間は「4_講師用進行台本」の所要時間から自動計算。学生の行動を「選ぶ→考える→発見する→提案する→顧客の反応を得る」へ段階的に変える')
hdr = ['フェーズ', '日本語', '予定（分）', '開始', '終了', '台本上の所要（分）', '差', 'スライド', '学生がすること', '参加型ワーク']
work = {
    'MISSION': '事前チェック（S02）',
    'HEARING': '3択×2（S06・S10）。考える30秒＋理由を聞く',
    'ANALYSIS': '自由記述：仮説を立てる（S16）／3択（S20）／グループ：原因を決めつけずに調べるには？（S22）',
    'DISCOVERY': '自由記述：この一言の裏は？（S25）／中心ワーク：誤カウントがなくなると何がうれしい？（S27）／3段目を言葉にする（S28）',
    'PROPOSAL': '模擬提案ワーク（S35：考える2分＋まとめる3分）／提案タイム（S36：90秒×2組＋顧客役の反応）',
    'REFLECTION': 'ペアで振り返り（S41）／事後アンケート（S45）',
}
rows = []
for i, p in enumerate(C['PHASES'], 5):
    ids = [s['id'] for s in C['SLIDES'] if s['phase'] == p['key']]
    rows.append([p['key'], p['ja'], p['end'] - p['start'], p['start'], p['end'], f"=SUMIF('4_講師用進行台本'!$B$5:$B${last4},A{i},'4_講師用進行台本'!$D$5:$D${last4})", f'=F{i}-C{i}', f"{ids[0]}〜{ids[-1]}（{len(ids)}枚）", p['lead'], work[p['key']]])
last2 = table(ws, 4, hdr, rows, [14, 10, 10, 7, 7, 14, 7, 18, 34, 60])
ws.cell(row=last2 + 1, column=1, value='合計').font = Font(name=FONT, size=10, bold=True)
for col, f in ((3, f'=SUM(C5:C{last2})'), (6, f'=SUM(F5:F{last2})'), (7, f'=SUM(G5:G{last2})')):
    ws.cell(row=last2 + 1, column=col, value=f).font = Font(name=FONT, size=10, bold=True)
ws.cell(row=last2 + 1, column=8, value=f'=COUNTA(\'4_講師用進行台本\'!A5:A{last4})&"枚"').font = Font(name=FONT, size=10, bold=True)
r = last2 + 3
notes = [
    '時間の内訳（目安）：動画 約3分／学生が考える・話す・発表する 約22分／講師の説明 約35分',
    'PROPOSALの模擬提案ワークは、指示書の推奨（3分＋4分＋4分＋2分）より短い「考える2分＋まとめる3分＋発表と反応4分」。理由と代案は「8_要確認事項」',
    '時間を超えそうなときの削り方：S19（仮説のイメージ）→ S44（若手社員：素材がないうちは非表示）→ S26右側 → S09を一言に。合計で約5〜7分短くできる',
    '発表の組数を増やすとき：1組増えるごとに約2分。上の削り方で時間をつくる',
]
for k, t in enumerate(notes):
    c = ws.cell(row=r + k, column=1, value=f'・{t}')
    c.font = Font(name=FONT, size=10)

# ── ⑤ 画像・動画差し替え管理 ──
ws = sheet('5_画像動画差し替え管理', '成果物⑤　画像・動画差し替え管理表', 'スライド上の枠は、PowerPointの「選択ウィンドウ」で同じ番号の名前がついている。枠の比率に合わせて用意すると、余白なく収まる')
hdr = ['番号', '種類', '用途', '配置フェーズ', 'スライド', '目的', '枠の大きさ（インチ）', '枠の比率', '推奨の大きさ（px）', '仮素材', '長さの上限', '状況', '差し替え予定日', '担当', '備考']
rows = []
for p in PLACED:
    m = media_of[p['id']]
    ratio = p['w'] / p['h']
    named = {16 / 9: '16:9', 4 / 3: '4:3', 1: '1:1', 3 / 4: '3:4'}
    near = min(named, key=lambda k: abs(k - ratio))
    rlabel = named[near] if abs(near - ratio) < 0.06 else f'{ratio:.2f}:1'
    rows.append([m['id'], m['type'], m['use'], m['phase'], p['slide'], m['purpose'], f"{p['w']}×{p['h']}", rlabel, f"{round(p['w'] * 300)}×{round(p['h'] * 300)}以上", m['tmp'] if p['img'] else 'なし（点線の枠のみ）', f"{m['maxSec']}秒以内" if m.get('maxSec') else '', '仮画像を配置' if p['img'] else '未提供', '', '', ''])
last5 = table(ws, 4, hdr, rows, [24, 7, 30, 16, 8, 36, 14, 10, 16, 36, 10, 14, 14, 10, 24])
dv = DataValidation(type='list', formula1='"未提供,仮画像を配置,素材受領,差し替え済み,確認済み"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f'L5:L{last5}')
for i in range(5, last5 + 1):
    for col in (12, 13, 14, 15):
        ws.cell(row=i, column=col).fill = PatternFill('solid', fgColor=YELLOW)
r = last5 + 2
rules = [
    '差し替えのルール（指示書14章）',
    '1. 指定画像を最優先し、仮画像を置き換える。被写体・形状・重要なディテールは加工しない',
    '2. 縦横比を保つ（引き伸ばさない）。比率が枠と違うときは、枠の中に収めて余白を背景色で埋める',
    '3. 背景透過画像は透明部分を保つ',
    '4. 既存デザインとの統一感は、背景・余白・文字配置で調整する',
    '5. 動画は「長さの上限」以内。上限を超える場合は「4_講師用進行台本」の所要時間を見直す',
    '6. 追加の画像・動画は IMAGE_PLACEHOLDER_12／VIDEO_PLACEHOLDER_06 から連番で、この表に行を足す',
]
for k, t in enumerate(rules):
    c = ws.cell(row=r + k, column=1, value=t)
    c.font = Font(name=FONT, size=10, bold=k == 0)
ws.cell(row=r + len(rules) + 1, column=1, value='状況ごとの件数').font = Font(name=FONT, size=10, bold=True)
for k, lab in enumerate(['未提供', '仮画像を配置', '素材受領', '差し替え済み', '確認済み']):
    ws.cell(row=r + len(rules) + 2 + k, column=1, value=lab).font = Font(name=FONT, size=10)
    ws.cell(row=r + len(rules) + 2 + k, column=2, value=f'=COUNTIF($L$5:$L${last5},A{r + len(rules) + 2 + k})').font = Font(name=FONT, size=10, bold=True)

# ── ⑥ 品質確認 ──
ws = sheet('6_品質確認', '成果物⑥　品質確認', '○＝満たす　△＝条件付き・素材待ち')
last6 = table(ws, 4, ['確認項目', '確認した内容', '判定', '残る課題・前提'], C['QA'], [30, 80, 14, 50], {3: RED})
for i in range(5, last6 + 1):
    ws.cell(row=i, column=3).alignment = CENTER
    ws.cell(row=i, column=3).font = Font(name=FONT, size=11, bold=True, color=INK)
r = last6 + 2
extra = [
    '技術面で直した・明記した点',
    '・「ボトルの接近」は、ボトルがくっつくと区切りがなくなり、むしろ少なく数える方向に働く。増える原因かは要確認と明記した（S17・S18）',
    '・既存スライド67（間隔の変化）は、上から見た図でボトル同士が重なって描かれており、実際には起きない。本編から外し、補足X04で修正が必要と明記',
    '・KR-Qシリーズは、既設と同じ「回帰反射型」。違いは「透明体の検出に向けて設計されていること」とだけ表現し、性能の数値・他社との優劣は載せていない',
    '・既存の導入動画は、0:15〜1:35付近で「お客様が実現したいこと」「確認作業が増える→生産効率が下がる」を先に見せている。DISCOVERYの答えが先に見えるため、使う範囲を限定する',
    '・現状データ（＋0.5〜0.8％、最大＋3.0％、10,050〜10,080本／10,300本）は、1万本に対する計算と一致。スライドに「演習用の設定値」と明記',
]
for k, t in enumerate(extra):
    c = ws.cell(row=r + k, column=1, value=t)
    c.font = Font(name=FONT, size=10, bold=k == 0)

# ── ⑦ アンケート ──
ws = sheet('7_アンケート', '成果物⑦　成果測定（アンケート）', '事前（S02）と事後（S45）で同じ問いを聞き、変化を見る。採用成果は次回ステップへの参加率で確認する')
r = table(ws, 4, ['設問', '事前チェック（開始時・1分）', '回答形式'], C['SURVEY']['pre'], [10, 60, 50])
r = table(ws, r + 2, ['設問', '事後アンケート（終了時・2分）', '回答形式'], C['SURVEY']['post'], [10, 60, 50])
table(ws, r + 2, ['指標', '計算方法', '目標'], C['SURVEY']['kpi'], [10, 60, 50])
ws.freeze_panes = None

# ── ⑧ 要確認事項 ──
ws = sheet('8_要確認事項', '要確認事項', '推測で事実を補わず、確認が必要なものとして残した事項')
rows = [[i + 1, *t, '', '未確認'] for i, t in enumerate(C['TODO'])]
last8 = table(ws, 4, ['No', '分類', '内容', '関係するスライド・素材', '確認先（案）', '回答・決定事項', '状況'], rows, [5, 10, 80, 24, 16, 36, 10])
for i in range(5, last8 + 1):
    for col in (6, 7):
        ws.cell(row=i, column=col).fill = PatternFill('solid', fgColor=YELLOW)
dv = DataValidation(type='list', formula1='"未確認,確認中,解決"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f'G5:G{last8}')

for w in wb.worksheets:
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.orientation = 'landscape'
    w.page_setup.fitToHeight = 0
wb.save(OUT)
print(OUT)
